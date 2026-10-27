"""
pdf_to_word_converter.py
=========================
Pipeline: PDF -> Word (.docx) -> Teks Bersih (.txt) -> MySQL

Cara penggunaan:
    python Chatbot/pdf_to_word_converter.py                         # Konversi PDF -> Word saja
    python Chatbot/pdf_to_word_converter.py --extract               # + ekstrak teks ke .txt
    python Chatbot/pdf_to_word_converter.py --insert                # + dry-run parse & cek MySQL
    python Chatbot/pdf_to_word_converter.py --insert --execute      # + INSERT nyata ke MySQL
    python Chatbot/pdf_to_word_converter.py --surah 16              # Filter surah tertentu
    python Chatbot/pdf_to_word_converter.py --skip-convert --extract --insert --execute

Dependensi (install dulu):
    pip install pdf2docx python-docx sqlalchemy pymysql python-dotenv
"""

import os
import sys
import re
import argparse
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

# -----------------------------------------------------------------------
# Path Config
# -----------------------------------------------------------------------
WORKSPACE_ROOT = Path(__file__).resolve().parent.parent
SPLIT_PDF_DIR  = WORKSPACE_ROOT / "DATA" / "split pdf"
WORD_DIR       = WORKSPACE_ROOT / "DATA" / "word"
ENV_PATH       = WORKSPACE_ROOT / "Chatbot" / ".env"

from dotenv import load_dotenv
load_dotenv(ENV_PATH)
DB_URL = os.getenv("DATABASE_URL")

# -----------------------------------------------------------------------
# Surah info
# -----------------------------------------------------------------------
SURAH_INFO = {
    13: {"name": "Ar-Rad",   "total_ayat": 43},
    14: {"name": "Ibrahim",  "total_ayat": 52},
    15: {"name": "Al-Hijr",  "total_ayat": 99},
    16: {"name": "An-Nahl",  "total_ayat": 128},
    17: {"name": "Al-Isra",  "total_ayat": 111},
    18: {"name": "Al-Kahf",  "total_ayat": 110},
    19: {"name": "Maryam",   "total_ayat": 98},
}

# -----------------------------------------------------------------------
# Pola noise OCR
# -----------------------------------------------------------------------
HEADER_RE = re.compile(
    r"""(?x)^(?:
        Tafsir\s*Ath[\-\s]*[Tt]habari|TddwAt|Tatsir\s+Ath|TafurAthfl|
        TCstultt|Tds,f|Tq6tu|TafshNi|TafsbAth|
        SURAH\s+(?:HUUD|AR.RA.D|IBRAHIM|AL.HIJR|AN.NAHL|AL.ISRA|AL.KAHF|MARYAM)|
        Surah\s*(?:Huud|fluud|Hwd|Ar.Ra|Ibrahim|Al.Hijr|An.Nahl|Al.Isra|Al.Kahf|Maryam|An.Nrsr)|
        SurahHuud|SurrchHuud|SunclrHrlud|SurohHuud|SurahHtmd|
        SurahAn|SurchAn|SurahAl|SurchAl|
        \d{1,4}\s*
    )$""",
    re.IGNORECASE | re.MULTILINE,
)

FOOTNOTE_KW = [
    'ibnu abi hatim dalam tafsir', 'zad al masir', 'an-nukat wa al',
    "ma'alim at-tanzil", 'abdurrazzaq dalam tafsir', 'lihat tafsir',
    'ad-durr al-mantsur', 'ibnu katsir', 'fath al-bari', 'ibid.',
]

SURAH_PATTERNS = {
    13: r"Ar[\-\s]*Ra[''`]?d",
    14: r"Ibrahim",
    15: r"Al[\-\s]*Hijr",
    16: r"An[\-\s]*Nahl",
    17: r"Al[\-\s]*Isra[''`]?",
    18: r"Al[\-\s]*Kahf",
    19: r"Maryam",
}


# ======================================================================
# STEP 1: Konversi PDF -> Word (.docx)
# ======================================================================

def convert_pdf_to_word(surah_filter=None):
    try:
        from pdf2docx import Converter
    except ImportError:
        print("[!] pdf2docx belum terinstall. Jalankan: pip install pdf2docx")
        return False

    print("=" * 70)
    print("STEP 1: PDF -> Word (.docx)")
    print("=" * 70)

    total_ok = 0
    total_fail = 0

    for surah_dir in sorted(SPLIT_PDF_DIR.iterdir()):
        if not surah_dir.is_dir():
            continue
        try:
            surah_no = int(surah_dir.name)
        except ValueError:
            continue
        if surah_filter and surah_no != surah_filter:
            continue

        sname = SURAH_INFO.get(surah_no, {}).get("name", f"Surah {surah_no}")
        out_dir = WORD_DIR / str(surah_no)
        out_dir.mkdir(parents=True, exist_ok=True)

        pdfs = sorted(surah_dir.glob("*.pdf"))
        print(f"\n[Surah {surah_no}: {sname}] {len(pdfs)} file PDF")

        for pdf in pdfs:
            docx = out_dir / (pdf.stem + ".docx")
            if docx.exists():
                size_mb = docx.stat().st_size / 1024 / 1024
                print(f"  [SKIP] {pdf.name} -> sudah ada ({size_mb:.1f} MB)")
                continue

            print(f"  [...] Konversi {pdf.name} ...", end=" ", flush=True)
            try:
                cv = Converter(str(pdf))
                cv.convert(str(docx), start=0, end=None)
                cv.close()
                size_mb = docx.stat().st_size / 1024 / 1024
                print(f"[OK] {docx.name} ({size_mb:.1f} MB)")
                total_ok += 1
            except Exception as e:
                print(f"[GAGAL] {e}")
                total_fail += 1

    print(f"\n[STEP 1 SELESAI] Berhasil: {total_ok} | Gagal: {total_fail}")
    return total_fail == 0


# ======================================================================
# STEP 2: Word (.docx) -> Teks Bersih (.txt)
# ======================================================================

def clean_text(raw):
    raw = re.sub(r'(\w+)-\s*\n\s*(\w+)', r'\1\2', raw)
    raw = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', raw)
    raw = re.sub(r'[ \t]{2,}', ' ', raw)
    fixes = {
        'Atlatr': 'Allah', 'Allatr': 'Allah',
        'slahih': 'shahih', 'bafuwa': 'bahwa', 'adatah': 'adalah',
        'mercka': 'mereka',
    }
    for w, r in fixes.items():
        raw = raw.replace(w, r)
    return raw.strip()


def is_noise(line):
    s = line.strip()
    if not s:
        return False
    if HEADER_RE.match(s):
        return True
    if re.fullmatch(r'\d{1,4}', s):
        return True
    if any(k in s.lower() for k in FOOTNOTE_KW):
        return True
    if len(s) < 3:
        return True
    return False


def extract_docx_to_txt(docx_path):
    from docx import Document
    doc = Document(str(docx_path))
    lines = []
    for para in doc.paragraphs:
        raw = para.text
        if not raw.strip():
            continue
        cleaned = clean_text(raw)
        if cleaned and not is_noise(cleaned):
            lines.append(cleaned)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    raw = para.text.strip()
                    if not raw:
                        continue
                    cleaned = clean_text(raw)
                    if cleaned and not is_noise(cleaned):
                        lines.append(cleaned)
    return '\n'.join(lines)


def extract_texts_from_word(surah_filter=None):
    try:
        from docx import Document  # noqa
    except ImportError:
        print("[!] python-docx belum terinstall. Jalankan: pip install python-docx")
        return 0

    print("\n" + "=" * 70)
    print("STEP 2: Word (.docx) -> Teks Bersih (.txt)")
    print("=" * 70)

    total = 0
    for surah_dir in sorted(WORD_DIR.iterdir()):
        if not surah_dir.is_dir():
            continue
        try:
            surah_no = int(surah_dir.name)
        except ValueError:
            continue
        if surah_filter and surah_no != surah_filter:
            continue

        sname = SURAH_INFO.get(surah_no, {}).get("name", f"Surah {surah_no}")
        txt_dir = surah_dir / "txt"
        txt_dir.mkdir(exist_ok=True)

        docxs = sorted(surah_dir.glob("*.docx"))
        print(f"\n[Surah {surah_no}: {sname}] {len(docxs)} file Word")

        for docx in docxs:
            txt = txt_dir / (docx.stem + ".txt")
            print(f"  [...] {docx.name} -> txt/{txt.name}", end=" ", flush=True)
            try:
                text = extract_docx_to_txt(docx)
                if not text.strip():
                    print("[WARN] Teks kosong")
                    continue
                txt.write_text(text, encoding='utf-8')
                print(f"[OK] {len(text):,} chars")
                total += 1
            except Exception as e:
                print(f"[GAGAL] {e}")

    print(f"\n[STEP 2 SELESAI] Total .txt: {total}")
    return total


# ======================================================================
# STEP 3: Parse Ayat -> INSERT ke MySQL
# ======================================================================

def get_engine():
    if not DB_URL:
        print("[X] DATABASE_URL tidak ada di .env!")
        return None
    try:
        from sqlalchemy import create_engine
        return create_engine(
            DB_URL,
            pool_pre_ping=True,
            pool_size=5,
            max_overflow=10,
            connect_args={"connect_timeout": 30},
        )
    except Exception as e:
        print(f"[X] Koneksi database gagal: {e}")
        return None


def get_missing_verses(engine, surah_no):
    from sqlalchemy import text
    with engine.connect() as conn:
        rows = conn.execute(text("""
            SELECT v.number
            FROM verses v
            JOIN chapters c ON v.id_chapter = c.id
            LEFT JOIN verse_tafsirs vt ON vt.id_verse = v.id AND vt.id_tafsir = 7
            WHERE c.number_chapter = :s AND vt.id IS NULL
            ORDER BY v.number
        """), {"s": surah_no}).fetchall()
    return [r[0] for r in rows]


def get_verse_meta(engine, surah_no, ayat_no):
    from sqlalchemy import text
    with engine.connect() as conn:
        row = conn.execute(text("""
            SELECT v.id, v.text_uthmani
            FROM verses v
            JOIN chapters c ON v.id_chapter = c.id
            WHERE c.number_chapter = :s AND v.number = :a
        """), {"s": surah_no, "a": ayat_no}).fetchone()
    return (row[0], row[1]) if row else (None, "")


def parse_verses_from_text(full_text, surah_no, total_ayat):
    sname = SURAH_PATTERNS.get(surah_no, r"\w+")
    patt = re.compile(
        rf'\((?:QS|Qs)\.?\s*(?:{sname})[^)]*?[:|]\s*([0-9][0-9,\-\s]*)\)',
        re.IGNORECASE
    )
    anchors = []
    for m in patt.finditer(full_text):
        raw = m.group(1).strip()
        nums = []
        for part in re.split(r'[,\s\-]+', raw):
            if part.isdigit():
                n = int(part)
                if 1 <= n <= total_ayat:
                    nums.append(n)
        if nums:
            anchors.append((m.start(), nums))

    if not anchors:
        return {}

    result = {}
    for i, (pos, nums) in enumerate(anchors):
        end = anchors[i + 1][0] if i + 1 < len(anchors) else len(full_text)
        segment = full_text[pos:end].strip()
        for vnum in nums:
            if vnum not in result:
                result[vnum] = segment
            else:
                result[vnum] += "\n" + segment
    return result


def process_and_insert(surah_filter=None, dry_run=True):
    engine = get_engine()
    if engine is None:
        return

    mode = "DRY-RUN (simulasi)" if dry_run else "EXECUTE (INSERT ke MySQL)"
    print("\n" + "=" * 70)
    print(f"STEP 3: Parse dan Insert | {mode}")
    print("=" * 70)

    total_ok = total_skip = total_notfound = 0

    for surah_no, info in sorted(SURAH_INFO.items()):
        if surah_filter and surah_no != surah_filter:
            continue

        sname = info["name"]
        total_ayat = info["total_ayat"]
        txt_dir = WORD_DIR / str(surah_no) / "txt"

        if not txt_dir.exists() or not list(txt_dir.glob("*.txt")):
            print(f"\n[WARN] Surah {surah_no}: tidak ada .txt -> jalankan --extract dulu")
            continue

        missing = set(get_missing_verses(engine, surah_no))
        print(f"\n[Surah {surah_no}: {sname}]")
        print(f"  Ayat kosong DB : {sorted(missing) if missing else 'Tidak ada (lengkap)'}")

        if not missing:
            print("  [OK] Sudah lengkap.")
            continue

        combined = ""
        for tf in sorted(txt_dir.glob("*.txt")):
            combined += "\n\n" + tf.read_text(encoding='utf-8', errors='replace')

        verse_map = parse_verses_from_text(combined, surah_no, total_ayat)
        print(f"  Ayat terparsing: {len(verse_map)}")

        from sqlalchemy import text as sa_text

        for ayat_no in sorted(missing):
            tafsir = verse_map.get(ayat_no)
            if not tafsir:
                print(f"  [?]  Ayat {ayat_no:3d}: tidak ditemukan di teks Word")
                total_notfound += 1
                continue

            id_verse, arabic = get_verse_meta(engine, surah_no, ayat_no)
            if not id_verse:
                print(f"  [!]  Ayat {ayat_no:3d}: id_verse tidak ada di DB")
                total_skip += 1
                continue

            if len(tafsir.strip()) < 30:
                print(f"  [!!] Ayat {ayat_no:3d}: teks terlalu pendek ({len(tafsir)} chars), skip")
                total_skip += 1
                continue

            payload = f"\r\n{arabic}\r\n{tafsir}\r\n"

            if not dry_run:
                with engine.begin() as conn:
                    conn.execute(sa_text("""
                        INSERT INTO verse_tafsirs (id_verse, id_tafsir, text)
                        VALUES (:vid, 7, :txt)
                        ON DUPLICATE KEY UPDATE text = :txt
                    """), {"vid": id_verse, "txt": payload})

            status = "DRY" if dry_run else "INS"
            print(f"  [{status}] Ayat {ayat_no:3d} | {len(payload):,} chars | id_verse={id_verse}")
            total_ok += 1

    print(f"\n{'=' * 70}")
    print("RINGKASAN:")
    print(f"  Diproses/Insert     : {total_ok}")
    print(f"  Tidak ada di teks   : {total_notfound}")
    print(f"  Dilewati            : {total_skip}")
    if dry_run:
        print("\n  [INFO] Mode simulasi. Untuk INSERT nyata: --insert --execute")
    else:
        print("\n  [SUKSES] Data berhasil masuk ke verse_tafsirs!")
    print("=" * 70)


# ======================================================================
# MAIN
# ======================================================================

def main():
    parser = argparse.ArgumentParser(
        description="Pipeline PDF -> Word -> Teks Bersih -> MySQL (Tafsir Ath-Thabari)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Contoh:
  python Chatbot/pdf_to_word_converter.py
  python Chatbot/pdf_to_word_converter.py --extract
  python Chatbot/pdf_to_word_converter.py --insert
  python Chatbot/pdf_to_word_converter.py --insert --execute
  python Chatbot/pdf_to_word_converter.py --surah 16 --insert --execute
  python Chatbot/pdf_to_word_converter.py --skip-convert --extract --insert --execute
        """
    )
    parser.add_argument("--extract", action="store_true",
                        help="Ekstrak teks dari Word (.docx) ke .txt")
    parser.add_argument("--insert", action="store_true",
                        help="Parse .txt dan cek/insert ke MySQL (dry-run default)")
    parser.add_argument("--execute", action="store_true",
                        help="INSERT nyata ke MySQL (butuh --insert)")
    parser.add_argument("--surah", type=int, metavar="N",
                        help="Filter nomor surah (13-19)")
    parser.add_argument("--skip-convert", action="store_true",
                        help="Lewati konversi PDF->Word")
    args = parser.parse_args()

    dry_run = not args.execute

    print("\n" + "=" * 70)
    print("PIPELINE: PDF -> WORD -> TEKS BERSIH -> MYSQL")
    print(f"  Surah : {args.surah or 'Semua (13-19)'}")
    print(f"  Mode  : {'DRY-RUN' if dry_run else 'EXECUTE'}")
    print("=" * 70)

    if not args.skip_convert:
        convert_pdf_to_word(args.surah)
    else:
        print("\n[SKIP] Step 1 dilewati (--skip-convert)")

    if args.extract or args.insert:
        extract_texts_from_word(args.surah)

    if args.insert:
        process_and_insert(args.surah, dry_run=dry_run)
    else:
        print("\n[HINT] --extract  -> ekstrak teks Word ke .txt")
        print("[HINT] --insert   -> parse .txt dan cek ke MySQL (dry-run)")
        print("[HINT] --insert --execute -> INSERT nyata ke MySQL")


if __name__ == "__main__":
    main()